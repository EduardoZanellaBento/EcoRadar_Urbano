"""
Gerador do trabalho escrito da APS (EcoRadar Urbano) em Word (.docx).

Uso (a partir da raiz do projeto ou de documentacao/):
    documentacao/.venv/Scripts/python documentacao/gerar_aps.py          # .docx + PDF + verificação
    documentacao/.venv/Scripts/python documentacao/gerar_aps.py --sem-pdf # só o .docx

O script é reexecutável: sempre regenera tudo do zero.
  1. renderiza os diagramas Mermaid (figuras/*.mmd -> PNG) se o PNG estiver ausente
     ou mais antigo que o .mmd (use --diagramas para forçar);
  2. monta o documento (conteúdo em conteudo_aps.py) com python-docx;
  3. converte para PDF com o LibreOffice (headless), descobre em que página começa cada
     título e preenche o sumário com esses números (2ª passagem);
  4. grava verificacao_paginas.md (páginas por seção x limites oficiais, figuras citadas).
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
FIGURAS = AQUI / "figuras"
NOME_BASE = "APS_CC_7o8o_2026_EcoRadar_Urbano"
DOCX = AQUI / f"{NOME_BASE}.docx"
PDF = AQUI / f"{NOME_BASE}.pdf"
FONTE = "Arial"
LARGURA_UTIL_CM = 16.0  # A4 (21 cm) - 2,5 cm - 2,5 cm

# Limites oficiais de páginas por seção primária (enunciado da APS)
LIMITES = {
    "1": ("Objetivo do trabalho", 1, 2),
    "2": ("Introdução", 2, 4),
    "3": ("Fundamentos das tecnologias para dispositivos móveis escolhidas", 4, 8),
    "4": ("Plano de desenvolvimento da aplicação", 5, 15),
    "5": ("Projeto (estrutura) do programa", 3, 8),
    "6": ("Relatório com as linhas de código", None, 10),
}


# =====================================================================================
# Blocos (o conteúdo é descrito como uma lista de blocos; a numeração de figuras,
# tabelas, listagens e seções é resolvida antes da renderização)
# =====================================================================================
@dataclass
class Bloco:
    tipo: str
    dados: dict = field(default_factory=dict)


class Conteudo:
    """API usada por conteudo_aps.py para descrever o documento."""

    def __init__(self) -> None:
        self.blocos: list[Bloco] = []

    # --- estrutura ---
    def capa(self, **dados):
        self.blocos.append(Bloco("capa", dados))

    def indice(self):
        self.blocos.append(Bloco("indice"))

    def inicio_numeracao(self):
        """Quebra de seção: a partir daqui as páginas mostram o número (canto superior direito)."""
        self.blocos.append(Bloco("secao_numerada"))

    def h1(self, texto, chave=None):
        self.blocos.append(Bloco("titulo", {"nivel": 1, "texto": texto, "chave": chave}))

    def h2(self, texto, chave=None):
        self.blocos.append(Bloco("titulo", {"nivel": 2, "texto": texto, "chave": chave}))

    def h3(self, texto, chave=None):
        self.blocos.append(Bloco("titulo", {"nivel": 3, "texto": texto, "chave": chave}))

    # --- texto ---
    def p(self, texto, destaque=False, recuo=True, alinhamento="justificado"):
        self.blocos.append(Bloco("paragrafo", {"texto": texto, "destaque": destaque, "recuo": recuo, "alinhamento": alinhamento}))

    def lista(self, itens, marcador="alinea", destaque=False):
        """marcador: 'alinea' (a), b), c)...) ou 'ponto' (•)."""
        self.blocos.append(Bloco("lista", {"itens": itens, "marcador": marcador, "destaque": destaque}))

    def quebra_pagina(self):
        self.blocos.append(Bloco("quebra"))

    # --- elementos numerados ---
    def tabela(self, chave, titulo, cabecalho, linhas, larguras, fonte=None, alinhar=None):
        self.blocos.append(Bloco("tabela", dict(chave=chave, titulo=titulo, cabecalho=cabecalho, linhas=linhas, larguras=larguras, fonte=fonte, alinhar=alinhar)))

    def figura(self, chave, titulo, caminho, largura_cm, fonte=None):
        self.blocos.append(Bloco("figura", dict(chave=chave, titulo=titulo, caminho=caminho, largura=largura_cm, fonte=fonte)))

    def figuras_lado_a_lado(self, chave, titulo, imagens, largura_cm=4.8, fonte=None):
        """imagens: lista de (caminho, rótulo) exibidas lado a lado numa tabela sem bordas.
        largura_cm pode ser um número (todas iguais) ou uma lista com a largura de cada imagem."""
        self.blocos.append(Bloco("figuras", dict(chave=chave, titulo=titulo, imagens=imagens, largura=largura_cm, fonte=fonte)))

    def codigo(self, chave, titulo, arquivo, intervalos, comentario="//"):
        self.blocos.append(Bloco("codigo", dict(chave=chave, titulo=titulo, arquivo=arquivo, intervalos=intervalos, comentario=comentario)))


# =====================================================================================
# Formatação (python-docx + XML)
# =====================================================================================
def _fontes_rfonts(elemento_rpr):
    rfonts = elemento_rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        elemento_rpr.insert(0, rfonts)
    for atributo in list(rfonts.attrib):
        if "Theme" in atributo or "theme" in atributo:
            del rfonts.attrib[atributo]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONTE)


def _estilo_base(estilo, negrito=False, italico=False):
    estilo.font.name = FONTE
    estilo.font.size = Pt(12)
    estilo.font.bold = negrito
    estilo.font.italic = italico
    estilo.font.color.rgb = RGBColor(0, 0, 0)
    rpr = estilo.element.get_or_add_rPr()
    _fontes_rfonts(rpr)
    cor = rpr.find(qn("w:color"))
    for atributo in list(cor.attrib):
        if "theme" in atributo.lower():
            del cor.attrib[atributo]
    pf = estilo.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    return estilo


def _novo_estilo(doc, nome, **kw):
    try:
        estilo = doc.styles[nome]
    except KeyError:
        estilo = doc.styles.add_style(nome, WD_STYLE_TYPE.PARAGRAPH)
    estilo.base_style = doc.styles["Normal"]
    return _estilo_base(estilo, **kw)


def configurar_estilos(doc):
    # Padrões do documento: Arial 12, espaçamento 1,5, sem espaço entre parágrafos
    padroes = doc.styles.element.find(qn("w:docDefaults"))
    rpr_padrao = padroes.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _fontes_rfonts(rpr_padrao)
    for sz in ("w:sz", "w:szCs"):
        el = rpr_padrao.find(qn(sz))
        if el is None:
            el = OxmlElement(sz)
            rpr_padrao.append(el)
        el.set(qn("w:val"), "24")
    ppr_padrao = padroes.find(qn("w:pPrDefault"))
    if ppr_padrao is not None and ppr_padrao.find(qn("w:pPr")) is not None:
        esp = ppr_padrao.find(qn("w:pPr")).find(qn("w:spacing"))
        if esp is not None:
            esp.set(qn("w:line"), "360")
            esp.set(qn("w:lineRule"), "auto")
            esp.set(qn("w:after"), "0")

    # Todos os estilos do modelo (inclusive os não usados, como Title e Heading 4-9) em Arial 12 / 1,5,
    # para que nenhum estilo aplicado depois no Word fuja da formatação exigida
    for estilo in doc.styles:
        tipo = estilo.type
        if tipo in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            estilo.font.name = FONTE
            estilo.font.size = Pt(12)
            _fontes_rfonts(estilo.element.get_or_add_rPr())
        if tipo == WD_STYLE_TYPE.PARAGRAPH:
            estilo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    normal = _estilo_base(doc.styles["Normal"])
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.widow_control = True

    for nivel in (1, 2, 3):
        h = doc.styles[f"Heading {nivel}"]
        _estilo_base(h, negrito=True, italico=(nivel == 3))
        pf = h.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.first_line_indent = Cm(0)
        pf.keep_with_next = True
        pf.space_before = Pt(0 if nivel == 1 else 12)
        pf.space_after = Pt(12 if nivel == 1 else 6)
        pf.page_break_before = nivel == 1
        if nivel == 1:
            h.font.all_caps = True

    s = _novo_estilo(doc, "APS Capa")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.first_line_indent = Cm(0)

    s = _novo_estilo(doc, "APS Titulo sem numero", negrito=True)
    s.font.all_caps = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.space_after = Pt(12)

    s = _novo_estilo(doc, "APS Legenda")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.space_before = Pt(6)

    s = _novo_estilo(doc, "APS Fonte")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.space_after = Pt(6)

    s = _novo_estilo(doc, "APS Figura")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE  # só a imagem (sem texto)

    s = _novo_estilo(doc, "APS Tabela")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.first_line_indent = Cm(0)

    s = _novo_estilo(doc, "APS Codigo")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.first_line_indent = Cm(0)
    s.paragraph_format.left_indent = Cm(0.5)
    s.paragraph_format.right_indent = Cm(0.2)
    s.paragraph_format.keep_together = True
    _sombreamento_paragrafo(s.element.get_or_add_pPr(), "EDEDED")

    s = _novo_estilo(doc, "APS Lista")
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    s.paragraph_format.left_indent = Cm(1.9)
    s.paragraph_format.first_line_indent = Cm(-0.65)

    for nivel, recuo in ((1, 0), (2, 0.6), (3, 1.2)):
        s = _novo_estilo(doc, f"toc {nivel}", negrito=(nivel == 1))
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.first_line_indent = Cm(0)
        s.paragraph_format.left_indent = Cm(recuo)
        s.paragraph_format.right_indent = Cm(0.8)
        s.paragraph_format.tab_stops.add_tab_stop(Cm(LARGURA_UTIL_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if nivel == 1:
            s.font.all_caps = True
        # o número da página pode ultrapassar o recuo direito (efeito "hanging" do sumário)
        ppr = s.element.get_or_add_pPr()
        ind = ppr.find(qn("w:ind"))
        ind.set(qn("w:right"), str(int(Cm(0.8).twips)))


def _sombreamento_paragrafo(ppr, cor):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    ppr.append(shd)


def _sombreamento_celula(celula, cor):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def _campo(paragrafo, instrucao, resultado="1"):
    def fld(tipo):
        r = paragrafo.add_run()
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), tipo)
        r._r.append(f)
        return r

    fld("begin")
    r = paragrafo.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instrucao} "
    r._r.append(it)
    fld("separate")
    paragrafo.add_run(resultado)
    fld("end")


def _fld_char(paragrafo, tipo):
    r = paragrafo.add_run()
    f = OxmlElement("w:fldChar")
    f.set(qn("w:fldCharType"), tipo)
    r._r.append(f)


# Marcação em linha: **negrito**, __itálico__, ==destaque amarelo==, `código` (texto literal,
# sem formatação extra: todo o documento é Arial 12). [REVISAR...]/[CONFERIR...] são destacados.
_PADRAO_INLINE = re.compile(r"(`[^`]+`|\*\*.+?\*\*|__.+?__|==.+?==|\[(?:REVISAR|CONFERIR)[^\]]*\])")


def escrever_texto(paragrafo, texto, destaque=False, negrito=False):
    for parte in _PADRAO_INLINE.split(texto):
        if not parte:
            continue
        b, i, h = negrito, False, destaque
        if parte.startswith("`") and parte.endswith("`"):
            run = paragrafo.add_run(parte[1:-1])
            run.bold = b or None
            if h:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            continue
        if parte.startswith("**") and parte.endswith("**"):
            parte, b = parte[2:-2], True
        elif parte.startswith("__") and parte.endswith("__"):
            parte, i = parte[2:-2], True
        elif parte.startswith("==") and parte.endswith("=="):
            parte, h = parte[2:-2], True
        elif parte.startswith("[REVISAR") or parte.startswith("[CONFERIR"):
            h, b = True, True
        for sub in _PADRAO_INLINE.split(parte) if _PADRAO_INLINE.search(parte) else [parte]:
            if not sub:
                continue
            # marcação aninhada simples (ex.: ==**texto**==)
            bb, ii = b, i
            if sub.startswith("**") and sub.endswith("**"):
                sub, bb = sub[2:-2], True
            elif sub.startswith("__") and sub.endswith("__"):
                sub, ii = sub[2:-2], True
            run = paragrafo.add_run(sub)
            run.bold = bb or None
            run.italic = ii or None
            if h:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW


# =====================================================================================
# Renderização
# =====================================================================================
class Renderizador:
    def __init__(self, conteudo: Conteudo, paginas_titulos: dict | None):
        self.c = conteudo
        self.paginas = paginas_titulos or {}
        self.doc = Document()
        self.num = {"F": {}, "T": {}, "L": {}, "S": {}}
        self.titulos: list[tuple[int, str, str]] = []  # (nível, número, texto)
        self.citacoes: dict[str, set] = {"F": set(), "T": set(), "L": set()}

    # --- numeração ---
    def numerar(self):
        cont = {"F": 0, "T": 0, "L": 0}
        h = [0, 0, 0]
        for b in self.c.blocos:
            if b.tipo in ("figura", "figuras"):
                cont["F"] += 1
                self.num["F"][b.dados["chave"]] = cont["F"]
            elif b.tipo == "tabela":
                cont["T"] += 1
                self.num["T"][b.dados["chave"]] = cont["T"]
            elif b.tipo == "codigo":
                cont["L"] += 1
                self.num["L"][b.dados["chave"]] = cont["L"]
            elif b.tipo == "titulo":
                n = b.dados["nivel"]
                h[n - 1] += 1
                for k in range(n, 3):
                    h[k] = 0
                numero = ".".join(str(x) for x in h[:n])
                b.dados["numero"] = numero
                self.titulos.append((n, numero, b.dados["texto"]))
                if b.dados.get("chave"):
                    self.num["S"][b.dados["chave"]] = numero

    _REF = re.compile(r"\{([FTLS]):([a-z0-9_]+)\}")

    def resolver(self, texto):
        def troca(m):
            tipo, chave = m.group(1), m.group(2)
            if chave not in self.num[tipo]:
                raise KeyError(f"Referência inexistente: {m.group(0)}")
            if tipo in self.citacoes:
                self.citacoes[tipo].add(chave)
            n = self.num[tipo][chave]
            return {"F": f"Figura {n}", "T": f"Tabela {n}", "L": f"Listagem {n}", "S": f"{n}"}[tipo]

        return self._REF.sub(troca, texto)

    # --- documento ---
    def configurar(self):
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.5)
        sec.header_distance = Cm(1.25)
        sec.footer_distance = Cm(1.25)
        configurar_estilos(doc)
        propriedades = doc.core_properties
        propriedades.title = "EcoRadar Urbano: aplicação distribuída para gerenciamento de informações ambientais urbanas"
        propriedades.subject = "APS — Ciência da Computação 7º/8º semestre — UNIP 2026"
        propriedades.author = "Grupo APS [REVISAR]"
        propriedades.language = "pt-BR"

    def renderizar(self) -> Document:
        self.configurar()
        self.numerar()
        for b in self.c.blocos:
            getattr(self, f"_r_{b.tipo}")(b.dados)
        return self.doc

    # --- blocos ---
    def _r_capa(self, d):
        doc = self.doc

        def linha(texto="", negrito=False, destaque=False, caixa_alta=False):
            p = doc.add_paragraph(style="APS Capa")
            escrever_texto(p, texto.upper() if caixa_alta else texto, destaque=destaque, negrito=negrito)
            return p

        linha("UNIP — Universidade Paulista", negrito=True, caixa_alta=True)
        linha("Ciência da Computação", negrito=True, caixa_alta=True)
        linha("7º/8º semestre")
        for _ in range(3):
            linha()
        linha("Atividades Práticas Supervisionadas (APS)", negrito=True)
        linha("Tema: “" + d["tema"] + "”")
        linha()
        linha(d["titulo"], negrito=True)
        for _ in range(3):
            linha()
        linha("Integrantes do grupo:", negrito=True)
        for aluno in d["alunos"]:
            linha(aluno)
        linha()
        linha("Professor responsável: " + d["professor"])
        for _ in range(4):
            linha()
        linha(d["campus"])
        linha(d["cidade"])
        p = linha(d["ano"])
        p.add_run().add_break(WD_BREAK.PAGE)

    def _r_indice(self, _d):
        doc = self.doc
        doc.add_paragraph("Índice", style="APS Titulo sem numero")
        entradas = self.titulos
        for i, (nivel, numero, texto) in enumerate(entradas):
            par = doc.add_paragraph(style=f"toc {nivel}")
            if i == 0:
                _fld_char(par, "begin")
                r = par.add_run()
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = ' TOC \\o "1-3" \\h \\z \\u '
                r._r.append(it)
                _fld_char(par, "separate")
            pagina = self.paginas.get(numero, "00")
            par.add_run(f"{numero} {texto}\t{pagina}")
            if i == len(entradas) - 1:
                _fld_char(par, "end")

    def _r_secao_numerada(self, _d):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.header.is_linked_to_previous = False
        cab = sec.header.paragraphs[0]
        cab.style = self.doc.styles["APS Tabela"]
        cab.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _campo(cab, "PAGE")
        self._logo_apos_secao = True

    def _r_titulo(self, d):
        texto = f'{d["numero"]} {d["texto"]}'
        p = self.doc.add_paragraph(style=f'Heading {d["nivel"]}')
        escrever_texto(p, texto)
        if d["nivel"] == 1 and getattr(self, "_logo_apos_secao", False):
            # a quebra de seção já inicia uma nova página
            p.paragraph_format.page_break_before = False
            self._logo_apos_secao = False

    def _r_paragrafo(self, d):
        p = self.doc.add_paragraph(style="Normal")
        if not d["recuo"]:
            p.paragraph_format.first_line_indent = Cm(0)
        if d["alinhamento"] == "centro":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif d["alinhamento"] == "esquerda":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(6)
        escrever_texto(p, self.resolver(d["texto"]), destaque=d["destaque"])

    def _r_lista(self, d):
        itens = d["itens"]
        for i, item in enumerate(itens):
            p = self.doc.add_paragraph(style="APS Lista")
            prefixo = f"{chr(ord('a') + i)})" if d["marcador"] == "alinea" else "•"
            p.paragraph_format.tab_stops.add_tab_stop(Cm(1.9))
            escrever_texto(p, f"{prefixo}\t" + self.resolver(item), destaque=d["destaque"])

    def _r_quebra(self, _d):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _legenda(self, texto):
        p = self.doc.add_paragraph(style="APS Legenda")
        escrever_texto(p, texto)
        return p

    def _fonte(self, texto=None):
        p = self.doc.add_paragraph(style="APS Fonte")
        escrever_texto(p, texto or "Fonte: elaborado pelos autores (2026).")
        return p

    def _r_figura(self, d):
        n = self.num["F"][d["chave"]]
        self._legenda(f"Figura {n} – " + self.resolver(d["titulo"]))
        p = self.doc.add_paragraph(style="APS Figura")
        p.add_run().add_picture(preparar_imagem(d["caminho"], d["largura"]), width=Cm(d["largura"]))
        self._fonte(d["fonte"])

    def _r_figuras(self, d):
        n = self.num["F"][d["chave"]]
        self._legenda(f"Figura {n} – " + self.resolver(d["titulo"]))
        imagens = d["imagens"]
        larguras = d["largura"] if isinstance(d["largura"], (list, tuple)) else [d["largura"]] * len(imagens)
        tabela = self.doc.add_table(rows=2, cols=len(imagens))
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabela.autofit = False
        margem = 0.45  # margens internas padrão da célula (2 x 0,19 cm) + folga
        total = sum(larguras) + margem * len(imagens)
        if total > LARGURA_UTIL_CM + 0.01:
            raise ValueError(f"Figuras lado a lado ({d['chave']}) com {total:.2f} cm: excede a largura útil")
        folga = (LARGURA_UTIL_CM - total) / len(imagens)
        _definir_larguras(tabela, [x + margem + folga for x in larguras])
        for j, ((caminho, rotulo), largura) in enumerate(zip(imagens, larguras)):
            celula = tabela.cell(0, j)
            p = celula.paragraphs[0]
            p.style = self.doc.styles["APS Figura"]
            p.add_run().add_picture(preparar_imagem(caminho, largura), width=Cm(largura))
            q = tabela.cell(1, j).paragraphs[0]
            q.style = self.doc.styles["APS Tabela"]
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            escrever_texto(q, f"({chr(ord('a') + j)}) {rotulo}")
        _impedir_quebra_linhas(tabela)
        self._fonte(d["fonte"])

    def _r_tabela(self, d):
        n = self.num["T"][d["chave"]]
        self._legenda(f"Tabela {n} – " + self.resolver(d["titulo"]))
        cab, linhas, larguras = d["cabecalho"], d["linhas"], d["larguras"]
        tabela = self.doc.add_table(rows=1 + len(linhas), cols=len(cab))
        tabela.style = self.doc.styles["Table Grid"]
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabela.autofit = False
        alinhar = d.get("alinhar") or ["esq"] * len(cab)
        for j, titulo in enumerate(cab):
            c = tabela.cell(0, j)
            _sombreamento_celula(c, "D9D9D9")
            p = c.paragraphs[0]
            p.style = self.doc.styles["APS Tabela"]
            escrever_texto(p, titulo, negrito=True)
        for i, linha in enumerate(linhas, start=1):
            for j, valor in enumerate(linha):
                p = tabela.cell(i, j).paragraphs[0]
                p.style = self.doc.styles["APS Tabela"]
                if alinhar[j] == "dir":
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif alinhar[j] == "centro":
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                escrever_texto(p, self.resolver(str(valor)))
        if abs(sum(larguras) - LARGURA_UTIL_CM) > 0.05:
            raise ValueError(f"Tabela {d['chave']}: larguras somam {sum(larguras):.2f} cm (esperado {LARGURA_UTIL_CM})")
        _definir_larguras(tabela, larguras)
        # cabeçalho repetido em cada página
        trpr = tabela.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        _impedir_quebra_linhas(tabela)
        self._fonte(d["fonte"])

    def _r_codigo(self, d):
        n = self.num["L"][d["chave"]]
        caminho = RAIZ / d["arquivo"]
        linhas = caminho.read_text(encoding="utf-8").splitlines()
        partes = []
        for k, (ini, fim) in enumerate(d["intervalos"]):
            if k:
                partes.append(None)
            partes.extend(linhas[ini - 1 : fim])
        textos = [x for x in partes if x is not None and x.strip()]
        recuo = min((len(x) - len(x.lstrip(" ")) for x in textos), default=0)
        faixa = " e ".join(f"{a}–{b}" if a != b else f"{a}" for a, b in d["intervalos"])
        rotulo = "linha" if len(d["intervalos"]) == 1 and d["intervalos"][0][0] == d["intervalos"][0][1] else "linhas"
        self._legenda(f"Listagem {n} – {self.resolver(d['titulo'])} ({d['arquivo']}, {rotulo} {faixa})")
        curta = len(partes) <= 10  # só listagens bem curtas ficam inteiras na mesma página
        for k, x in enumerate(partes):
            p = self.doc.add_paragraph(style="APS Codigo")
            if x is None:
                texto = f"{d['comentario']} (…)"
            else:
                texto = x[recuo:] if len(x) >= recuo else x.lstrip(" ")
            run = p.add_run(texto if texto else " ")
            run.font.name = FONTE
            if curta and k < len(partes) - 1:
                p.paragraph_format.keep_with_next = True
        # espaço após a listagem
        self.doc.paragraphs[-1].paragraph_format.space_after = Pt(6)


def _definir_larguras(tabela, larguras_cm):
    """Largura fixa por coluna: tblGrid/gridCol (usado pelo LibreOffice), tcW (Word) e tblW."""
    tbl = tabela._tbl
    grade = tbl.tblGrid
    for col in list(grade):
        grade.remove(col)
    for largura in larguras_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(Cm(largura).twips)))
        grade.append(gc)
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        existente = tblpr.find(qn(tag))
        if existente is not None:
            tblpr.remove(existente)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(int(Cm(sum(larguras_cm)).twips)))
    tblw.set(qn("w:type"), "dxa")
    tblpr.append(tblw)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras_cm):
            celula.width = Cm(largura)


CACHE_IMAGENS = AQUI / ".cache_imagens"
DPI_IMPRESSAO = 250


def preparar_imagem(caminho, largura_cm) -> str:
    """Reduz a imagem para ~250 dpi no tamanho impresso (o .docx não precisa da resolução original).
    Prints de tela viram JPEG de alta qualidade; diagramas continuam PNG. Os originais não são alterados."""
    from PIL import Image

    caminho = Path(caminho)
    alvo_px = int(largura_cm / 2.54 * DPI_IMPRESSAO)
    with Image.open(caminho) as im:
        if im.width <= alvo_px * 1.1 and caminho.stat().st_size < 400_000:
            return str(caminho)
        diagrama = caminho.parent == FIGURAS
        destino = CACHE_IMAGENS / f"{caminho.parent.name}_{caminho.stem}_{alvo_px}.{'png' if diagrama else 'jpg'}"
        if destino.exists() and destino.stat().st_mtime >= caminho.stat().st_mtime:
            return str(destino)
        CACHE_IMAGENS.mkdir(exist_ok=True)
        novo = im.convert("RGB")
        if novo.width > alvo_px:
            novo = novo.resize((alvo_px, round(novo.height * alvo_px / novo.width)), Image.LANCZOS)
        if diagrama:
            novo.save(destino, optimize=True)
        else:
            novo.save(destino, quality=90, optimize=True)
    return str(destino)


def _impedir_quebra_linhas(tabela):
    for linha in tabela.rows:
        trpr = linha._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trpr.append(cs)


# =====================================================================================
# LibreOffice: conversão, páginas dos títulos e verificação
# =====================================================================================
def localizar_soffice() -> str | None:
    candidatos = [
        shutil.which("soffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ]
    return next((c for c in candidatos if c and Path(c).exists()), None)


def converter_pdf(soffice: str) -> Path:
    subprocess.run([soffice, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", str(AQUI), str(DOCX)],
                   check=True, capture_output=True, timeout=300)
    return PDF


def _normalizar(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().upper()


def paginas_dos_titulos(pdf: Path, titulos) -> tuple[dict, int, list[str]]:
    import pymupdf

    doc = pymupdf.open(str(pdf))
    textos = [pg.get_text() for pg in doc]
    total = len(textos)
    # Páginas do índice: muitas linhas pontilhadas
    paginas_indice = {i for i, t in enumerate(textos) if t.count("....") >= 5}
    ultima_indice = max(paginas_indice) if paginas_indice else 1
    resultado, pagina_atual = {}, ultima_indice + 1
    for nivel, numero, texto in titulos:
        alvo = _normalizar(f"{numero} {texto}")[:45]
        encontrado = None
        for i in range(pagina_atual, total):
            linhas = [_normalizar(l) for l in textos[i].splitlines()]
            juntas = [linhas[k] + " " + linhas[k + 1] if k + 1 < len(linhas) else linhas[k] for k in range(len(linhas))]
            if any(l.startswith(alvo) for l in linhas) or any(j.startswith(alvo) for j in juntas):
                encontrado = i
                break
        if encontrado is None:
            raise RuntimeError(f"Título não encontrado no PDF: {numero} {texto}")
        resultado[numero] = encontrado + 1
        pagina_atual = encontrado
    return resultado, total, textos


def gravar_verificacao(paginas: dict, total: int, render: Renderizador, textos: list[str]):
    h1 = [(n, t) for (nv, n, t) in render.titulos if nv == 1]
    linhas = [
        "# Verificação de páginas — trabalho escrito da APS",
        "",
        f"Gerado por `gerar_aps.py` a partir do PDF convertido pelo LibreOffice (`{PDF.name}`, {total} páginas).",
        "A página inicial de cada seção foi localizada extraindo o texto de cada página do PDF (PyMuPDF);",
        "como toda seção primária começa em página nova, páginas da seção = início da próxima − início da seção.",
        "",
        "| Seção | Páginas no PDF | Nº de páginas | Limite oficial | Situação |",
        "|---|---|---|---|---|",
    ]
    todas_ok = True
    for idx, (numero, texto) in enumerate(h1):
        ini = paginas[numero]
        fim = (paginas[h1[idx + 1][0]] - 1) if idx + 1 < len(h1) else total
        qtd = fim - ini + 1
        limite = LIMITES.get(numero)
        if limite:
            _, minimo, maximo = limite
            ok = (minimo is None or qtd >= minimo) and (maximo is None or qtd <= maximo)
            todas_ok &= ok
            faixa = f"{minimo if minimo else '—'} a {maximo}"
            situacao = "OK" if ok else "**ajustar**"
        else:
            faixa, situacao = "sem limite", "—"
        linhas.append(f"| {numero} {texto} | {ini}–{fim} | {qtd} | {faixa} | {situacao} |")
    linhas += ["", f"**Resultado:** {'todas as seções dentro dos limites.' if todas_ok else 'há seções fora dos limites.'}", ""]

    # Figuras/tabelas/listagens: existentes x citadas
    linhas += ["## Figuras, tabelas e listagens", "", "| Tipo | Inseridas | Citadas no texto | Não citadas |", "|---|---|---|---|"]
    for tipo, nome in (("F", "Figuras"), ("T", "Tabelas"), ("L", "Listagens")):
        inseridas = set(render.num[tipo])
        citadas = render.citacoes[tipo]
        faltando = sorted(inseridas - citadas, key=lambda k: render.num[tipo][k])
        linhas.append(f"| {nome} | {len(inseridas)} | {len(citadas & inseridas)} | {', '.join(faltando) or 'nenhuma'} |")
    linhas += ["", "Referências a figuras/tabelas inexistentes interrompem a geração (erro), portanto todas as citadas existem.", ""]

    # Arquivos de imagem usados
    faltam = [str(b.dados.get("caminho")) for b in render.c.blocos if b.tipo == "figura" and not Path(b.dados["caminho"]).exists()]
    linhas += [f"Arquivos de imagem ausentes: {', '.join(faltam) if faltam else 'nenhum'}.", ""]

    linhas += auditoria_formatacao()

    # Marcadores pendentes
    completo = "\n".join(textos)
    pend = sorted({re.sub(r"\s+", " ", m) for m in re.findall(r"\[(?:REVISAR|CONFERIR|NOME COMPLETO \d|RA\d|CAMPUS|CIDADE)[^\]]*\]", completo)})
    linhas += ["## Marcadores que o grupo precisa revisar (encontrados no PDF)", ""]
    linhas += [f"- `{m}`" for m in pend] + [""]
    (AQUI / "verificacao_paginas.md").write_text("\n".join(linhas), encoding="utf-8")
    return todas_ok


def auditoria_formatacao() -> list[str]:
    """Confere fonte, tamanho, espaçamento e margens no .docx e no PDF gerados."""
    import zipfile

    import pymupdf

    xml = zipfile.ZipFile(DOCX).read("word/document.xml").decode("utf-8")
    estilos = zipfile.ZipFile(DOCX).read("word/styles.xml").decode("utf-8")
    tamanhos = sorted(set(re.findall(r'<w:sz w:val="(\d+)"', xml + estilos)))
    fontes_docx = sorted(set(re.findall(r'w:ascii="([^"]+)"', xml + estilos)))
    # espaçamento: 1,5 = w:line="360" (auto); a única exceção são os parágrafos que contêm só imagem
    espacamentos = sorted(set(re.findall(r'<w:spacing [^>]*w:line="(\d+)"', xml + estilos)))

    pdf = pymupdf.open(str(PDF))
    fontes_pdf, fora = set(), []
    cm = 72 / 2.54
    for n, pg in enumerate(pdf, start=1):
        for f in pg.get_fonts():
            fontes_pdf.add(f[3].split("+")[-1])
        for img in pg.get_image_info():
            x0, y0, x1, y1 = img["bbox"]
            if x0 < 2.5 * cm - 3 or x1 > (21 - 2.5) * cm + 3 or y0 < 2.5 * cm - 3 or y1 > (29.7 - 2.5) * cm + 3:  # tolerância de 3 pt (~1 mm)
                fora.append(f"p. {n}")
    tamanhos_pdf = set()
    for pg in pdf:
        for bloco in pg.get_text("dict")["blocks"]:
            for linha in bloco.get("lines", []):
                for s in linha["spans"]:
                    if s["text"].strip():
                        tamanhos_pdf.add(round(s["size"], 1))
    return [
        "## Auditoria de formatação",
        "",
        f"- Tamanhos de fonte declarados no .docx (meios-pontos): {', '.join(tamanhos)} → "
        f"{'somente 12 pt' if tamanhos == ['24'] else 'ATENÇÃO: há outros tamanhos'}",
        f"- Fontes declaradas no .docx: {', '.join(fontes_docx)}",
        f"- Espaçamentos de linha no .docx (w:line): {', '.join(espacamentos)} (360 = 1,5; 240 = simples, usado só nos parágrafos que contêm apenas imagem)",
        f"- Fontes embutidas no PDF: {', '.join(sorted(fontes_pdf))}",
        f"- Tamanhos de texto no PDF (pt): {', '.join(str(t) for t in sorted(tamanhos_pdf))}",
        f"- Imagens fora das margens de 2,5 cm: {', '.join(fora) if fora else 'nenhuma'}",
        "",
    ]


def renderizar_diagramas(forcar: bool):
    pendentes = [m for m in FIGURAS.glob("*.mmd") if forcar or not m.with_suffix(".png").exists()
                 or m.with_suffix(".png").stat().st_mtime < m.stat().st_mtime]
    if not pendentes:
        return
    node = shutil.which("node")
    if not node:
        print("AVISO: Node.js não encontrado; usando os PNG existentes dos diagramas.")
        return
    nomes = [m.stem for m in pendentes]
    subprocess.run([node, str(AQUI / "renderizar_diagramas.mjs"), *nomes], check=True)


def gerar(paginas=None) -> Renderizador:
    sys.path.insert(0, str(AQUI))
    import importlib

    import conteudo_aps

    importlib.reload(conteudo_aps)
    conteudo = Conteudo()
    conteudo_aps.montar(conteudo, RAIZ, AQUI)
    render = Renderizador(conteudo, paginas)
    doc = render.renderizar()
    doc.save(str(DOCX))
    return render


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--sem-pdf", action="store_true", help="gera apenas o .docx (sem LibreOffice)")
    ap.add_argument("--diagramas", action="store_true", help="força a renderização dos diagramas Mermaid")
    args = ap.parse_args()

    renderizar_diagramas(args.diagramas)
    render = gerar()
    print(f"1ª passagem: {DOCX.name} gerado")
    if args.sem_pdf:
        return
    soffice = localizar_soffice()
    if not soffice:
        print("LibreOffice não encontrado: o sumário ficará com páginas '00' — abra no Word e pressione F9.")
        return
    converter_pdf(soffice)
    paginas, total, _ = paginas_dos_titulos(PDF, render.titulos)
    # 2ª passagem: sumário com os números reais (o comprimento das linhas do sumário não muda)
    render = gerar(paginas)
    converter_pdf(soffice)
    paginas2, total2, textos = paginas_dos_titulos(PDF, render.titulos)
    if paginas2 != paginas:
        render = gerar(paginas2)
        converter_pdf(soffice)
        paginas2, total2, textos = paginas_dos_titulos(PDF, render.titulos)
    ok = gravar_verificacao(paginas2, total2, render, textos)
    print(f"PDF: {PDF.name} ({total2} páginas). Limites: {'OK' if ok else 'AJUSTAR'} — ver verificacao_paginas.md")


if __name__ == "__main__":
    main()
